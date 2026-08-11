from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import math
import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Mapping, Optional, Sequence, Tuple

from autotrainer.core.animal.external_metadata import (
    ExternalAnimalRecord,
    ExternalIdentity,
    NormalizedAnimalBatch,
    normalize_optional_text,
    normalize_rfid,
    normalize_state,
    split_genotype,
)


def _utc_now() -> str:
    return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace(
        "+00:00", "Z"
    )


def _header_key(value: Any) -> str:
    text = normalize_optional_text(value)
    return "" if text is None else re.sub(r"[^a-z0-9]", "", text.casefold())


def _json_value(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _identifier_value(value: Any) -> Optional[str]:
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_optional_text(value)


def _record_hash(payload: Mapping[str, Any]) -> str:
    encoded = json.dumps(
        payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


@dataclass(frozen=True)
class SoftMouseMappingProfile:
    profile_id: str = "softmouse-default"
    version: int = 2
    sheet_name: str = "Animal List"
    external_identity_column: str = "Physical Tag"
    rfid_column: str = "Plate ID"
    new_animal_name_column: str = "Physical Tag"
    state_column: str = "State"
    sex_column: Optional[str] = "Sex"
    date_of_birth_column: Optional[str] = "Date of Birth"
    strain_column: Optional[str] = "Strain"
    genotype_column: Optional[str] = "Genotype"
    cage_column: Optional[str] = "Cage Tag"
    protocol_column: Optional[str] = "Protocol"
    ended_states: Tuple[str, ...] = ("ended",)

    def required_columns(self) -> Tuple[str, ...]:
        return tuple(
            dict.fromkeys(
                (
                    self.external_identity_column,
                    self.rfid_column,
                    self.new_animal_name_column,
                    self.state_column,
                )
            )
        )

    def configured_columns(self) -> Tuple[str, ...]:
        optional = (
            self.sex_column,
            self.date_of_birth_column,
            self.strain_column,
            self.genotype_column,
            self.cage_column,
            self.protocol_column,
        )
        return self.required_columns() + tuple(item for item in optional if item)


@dataclass(frozen=True)
class ImportGuardrails:
    minimum_source_rows: int = 1
    maximum_fractional_row_drop: float = 0.35

    def validate(self, row_count: int, previous_row_count: Optional[int]) -> None:
        if row_count < self.minimum_source_rows:
            raise ValueError(
                f"Source has {row_count} rows; expected at least "
                f"{self.minimum_source_rows} for a complete export"
            )
        if previous_row_count and previous_row_count > 0:
            minimum = math.ceil(
                previous_row_count * (1 - self.maximum_fractional_row_drop)
            )
            if row_count < minimum:
                raise ValueError(
                    f"Source row count dropped from {previous_row_count} to {row_count}; "
                    "refusing a potentially filtered export"
                )


@dataclass(frozen=True)
class SpreadsheetImportPreview:
    batch: NormalizedAnimalBatch
    source_path: Path
    source_size: int
    source_modified_ns: int
    source_headers: Tuple[str, ...]
    mapped_columns: Tuple[str, ...]
    ignored_columns: Tuple[str, ...]


class SoftMouseSpreadsheetSource:
    def __init__(
        self,
        profile: SoftMouseMappingProfile = SoftMouseMappingProfile(),
        guardrails: ImportGuardrails = ImportGuardrails(),
    ):
        self.profile = profile
        self.guardrails = guardrails

    def preview(
        self,
        path: Path,
        *,
        previous_source_row_count: Optional[int] = None,
    ) -> SpreadsheetImportPreview:
        source_path = Path(path)
        stat = source_path.stat()
        source_bytes = source_path.read_bytes()
        source_hash = hashlib.sha256(source_bytes).hexdigest()
        headers, source_rows, sheet_name = self._read_rows(source_path)
        resolved = self._resolve_headers(headers)
        self.guardrails.validate(len(source_rows), previous_source_row_count)

        records: List[ExternalAnimalRecord] = []
        ignored_missing_rfid = 0
        ignored_ended = 0
        ended_states = {normalize_state(value) for value in self.profile.ended_states}
        for row_number, values in enumerate(source_rows, start=2):
            payload = {
                header: _json_value(values[index]) if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            state = normalize_state(self._value(values, resolved[self.profile.state_column]))
            if state in ended_states:
                ignored_ended += 1
                continue
            raw_rfid = self._value(values, resolved[self.profile.rfid_column])
            if normalize_optional_text(raw_rfid) is None:
                ignored_missing_rfid += 1
                continue
            if not isinstance(raw_rfid, str):
                raise ValueError(
                    f"RFID at source row {row_number} is stored as a numeric cell; "
                    "format the SoftMouse RFID export column as text to prevent "
                    "spreadsheet precision loss"
                )
            try:
                rfid = normalize_rfid(raw_rfid)
            except ValueError:
                # Plate ID may also be used for non-RFID identifiers. Only an
                # exact ISO 11784 15-digit value participates in the registry.
                ignored_missing_rfid += 1
                continue
            external_id = _identifier_value(
                self._value(values, resolved[self.profile.external_identity_column])
            )
            if external_id is None:
                raise ValueError(
                    f"RFID-bearing source row {row_number} has no "
                    f"{self.profile.external_identity_column!r}"
                )
            records.append(
                ExternalAnimalRecord(
                    identity=ExternalIdentity(external_id),
                    physical_rfid=rfid,
                    new_animal_name_candidate=normalize_optional_text(
                        self._value(values, resolved[self.profile.new_animal_name_column])
                    ),
                    state=state,
                    source_payload=payload,
                    source_hash=_record_hash(payload),
                    sex=self._optional(values, resolved, self.profile.sex_column),
                    date_of_birth=self._optional(
                        values, resolved, self.profile.date_of_birth_column
                    ),
                    strain=self._optional(values, resolved, self.profile.strain_column),
                    genotype=split_genotype(
                        self._optional(values, resolved, self.profile.genotype_column)
                    ),
                    cage=self._optional(values, resolved, self.profile.cage_column),
                    protocol=self._optional(
                        values, resolved, self.profile.protocol_column
                    ),
                )
            )

        header_signature = hashlib.sha256(
            json.dumps(headers, separators=(",", ":")).encode("utf-8")
        ).hexdigest()
        batch = NormalizedAnimalBatch(
            import_id=str(uuid.uuid4()),
            provider="softmouse",
            schema_version=1,
            mapping_profile_id=self.profile.profile_id,
            mapping_profile_version=self.profile.version,
            source_filename=source_path.name,
            source_sheet=sheet_name,
            source_file_sha256=source_hash,
            header_signature=header_signature,
            imported_utc=_utc_now(),
            total_source_rows=len(source_rows),
            ignored_missing_rfid_rows=ignored_missing_rfid,
            ignored_ended_rows=ignored_ended,
            records=tuple(records),
        )
        return SpreadsheetImportPreview(
            batch=batch,
            source_path=source_path,
            source_size=stat.st_size,
            source_modified_ns=stat.st_mtime_ns,
            source_headers=tuple(headers),
            mapped_columns=tuple(
                header for index, header in enumerate(headers) if index in resolved.values()
            ),
            ignored_columns=tuple(
                header for index, header in enumerate(headers) if header and index not in resolved.values()
            ),
        )

    def _read_rows(self, path: Path) -> Tuple[List[str], List[Tuple[Any, ...]], str]:
        suffix = path.suffix.casefold()
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            if not rows:
                raise ValueError("Spreadsheet is empty")
            data_rows = [tuple(r) for r in rows[1:] if any(str(value).strip() for value in r)]
            return [str(value).strip() for value in rows[0]], data_rows, ""
        if suffix != ".xlsx":
            raise ValueError("SoftMouse source must be an .xlsx or .csv file")
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to import .xlsx exports") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if self.profile.sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Missing worksheet {self.profile.sheet_name!r}; found "
                    f"{workbook.sheetnames!r}"
                )
            sheet = workbook[self.profile.sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            try:
                raw_headers = next(iterator)
            except StopIteration as exc:
                raise ValueError("Spreadsheet is empty") from exc
            headers = [normalize_optional_text(value) or "" for value in raw_headers]
            rows = [tuple(row) for row in iterator if any(value is not None for value in row)]
            return headers, rows, sheet.title
        finally:
            workbook.close()

    def _resolve_headers(self, headers: Sequence[str]) -> Dict[str, int]:
        aliases: Dict[str, List[int]] = {}
        for index, header in enumerate(headers):
            aliases.setdefault(_header_key(header), []).append(index)
        result: Dict[str, int] = {}
        required = set(self.profile.required_columns())
        for configured in self.profile.configured_columns():
            matches = aliases.get(_header_key(configured), [])
            if len(matches) > 1:
                raise ValueError(f"Ambiguous source column {configured!r}")
            if not matches:
                if configured in required:
                    raise ValueError(f"Missing required source column {configured!r}")
                continue
            result[configured] = matches[0]
        return result

    @staticmethod
    def _value(values: Sequence[Any], index: int) -> Any:
        return values[index] if index < len(values) else None

    def _optional(
        self,
        values: Sequence[Any],
        resolved: Mapping[str, int],
        column: Optional[str],
    ) -> Optional[str]:
        if not column or column not in resolved:
            return None
        value = self._value(values, resolved[column])
        if isinstance(value, dt.datetime):
            return value.date().isoformat()
        if isinstance(value, dt.date):
            return value.isoformat()
        return normalize_optional_text(value)
